from __future__ import annotations

from pathlib import Path
import textwrap
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "workdir" / "handover"
DOCX_PATH = OUT_DIR / "RagAgent_TED文档问答助手_交接文档.docx"
IMG_DIR = OUT_DIR / "images"

DEPLOY_PATH = "/share/home/huangli/Rag_demo"


PALETTE = {
    "ink": "17201f",
    "muted": "667370",
    "line": "d9e2df",
    "surface": "ffffff",
    "soft": "f6faf9",
    "soft2": "eef5f4",
    "accent": "0f766e",
    "accent2": "2563eb",
    "success": "147c4f",
    "warning": "a25b10",
    "danger": "b42318",
}


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def _font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font,
    fill: str,
    max_width: int,
    line_gap: int = 6,
) -> int:
    x, y = xy
    lines: list[str] = []
    for raw_line in text.splitlines():
        if not raw_line:
            lines.append("")
            continue
        current = ""
        for ch in raw_line:
            candidate = current + ch
            if draw.textlength(candidate, font=font) <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    line_height = font.size + line_gap
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height
    return y


def rounded(draw, box, fill, outline="#d9e2df", radius=12, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 850), "#f6faf9")
    d = ImageDraw.Draw(img)
    title = _font(38, True)
    h = _font(24, True)
    body = _font(21)
    small = _font(18)

    d.text((56, 40), "RagAgent TED 文档问答助手链路", font=title, fill=f"#{PALETTE['ink']}")
    d.text((58, 90), "只覆盖 TED 文档问答助手：/ragagent、/v1/rag/ask、/v1/rag/reindex、/v1/query/rewrite", font=small, fill=f"#{PALETTE['muted']}")

    boxes = [
        (70, 170, 320, 315, "前端入口", "/ragagent\nGradio 工作台\n查询 / 改写 / 提问"),
        (380, 170, 660, 315, "Query Rewrite", "保守型 / 激进型\n补充 TED 术语\n不改变用户意图"),
        (720, 170, 1010, 315, "问答 Agent", "RagQaAgent.ask()\n加载缓存资产\n拒答阈值保护"),
        (1070, 170, 1485, 315, "答案与证据", "LLM 严格基于 evidence 回答\n返回 answer、status、evidence"),
        (170, 435, 445, 610, "索引构建", "ResourceIndexer\nHTML 正文抽取\nMarkdown/文本标题分段\n表格转 Markdown"),
        (505, 435, 780, 610, "向量索引", "EmbeddingRetriever\nworkdir/vector_index\nLATEST 指向最新索引"),
        (840, 435, 1115, 610, "混合召回", "向量召回 + BM25-like\n候选去重合并\nrerank 可降级"),
        (1175, 435, 1450, 610, "反馈闭环", "qa_feedback.db\n历史对话\n有用/无用评价"),
    ]

    for x1, y1, x2, y2, head, text in boxes:
        fill = "#ffffff"
        outline = f"#{PALETTE['line']}"
        rounded(d, (x1, y1, x2, y2), fill, outline, 18, 2)
        d.text((x1 + 24, y1 + 20), head, font=h, fill=f"#{PALETTE['accent'] if y1 < 400 else PALETTE['accent2']}")
        draw_wrapped(d, (x1 + 24, y1 + 62), text, body, f"#{PALETTE['ink']}", x2 - x1 - 48, 8)

    arrows = [
        ((320, 242), (380, 242)),
        ((660, 242), (720, 242)),
        ((1010, 242), (1070, 242)),
        ((585, 315), (585, 435)),
        ((445, 522), (505, 522)),
        ((780, 522), (840, 522)),
        ((1115, 522), (1175, 522)),
        ((975, 435), (925, 315)),
    ]
    for start, end in arrows:
        d.line([start, end], fill=f"#{PALETTE['accent']}", width=4)
        ex, ey = end
        sx, sy = start
        if ex > sx:
            pts = [(ex, ey), (ex - 14, ey - 9), (ex - 14, ey + 9)]
        elif ex < sx:
            pts = [(ex, ey), (ex + 14, ey - 9), (ex + 14, ey + 9)]
        elif ey > sy:
            pts = [(ex, ey), (ex - 9, ey - 14), (ex + 9, ey - 14)]
        else:
            pts = [(ex, ey), (ex - 9, ey + 14), (ex + 9, ey + 14)]
        d.polygon(pts, fill=f"#{PALETTE['accent']}")

    rounded(d, (70, 685, 1485, 775), "#eef5f4", f"#{PALETTE['line']}", 16, 1)
    d.text((95, 710), "部署主路径", font=h, fill=f"#{PALETTE['ink']}")
    d.text((245, 713), DEPLOY_PATH, font=_font(25, True), fill=f"#{PALETTE['accent2']}")
    d.text((95, 747), "维护重点：Resource 文档更新后必须重建索引；模型、Embedding、Rerank 地址和密钥只放在内网 .env 中。", font=small, fill=f"#{PALETTE['muted']}")
    img.save(path)


def make_ui_overview(path: Path) -> None:
    """Reuse the actual workbench screenshot maintained with the user guide."""
    source = ROOT / "docs" / "assets" / "workbench-desktop.png"
    if not source.is_file():
        raise FileNotFoundError(f"Missing workbench screenshot: {source}")
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, path)


def make_evidence_table(path: Path) -> None:
    img = Image.new("RGB", (1500, 850), "#ffffff")
    d = ImageDraw.Draw(img)
    h1 = _font(31, True)
    h2 = _font(22, True)
    body = _font(19)
    small = _font(16)
    mono = ImageFont.truetype("C:/Windows/Fonts/consola.ttf", 16) if Path("C:/Windows/Fonts/consola.ttf").exists() else small

    d.text((30, 25), "Evidence 证据区示意：来源、分数、结构化表格", font=h1, fill=f"#{PALETTE['ink']}")
    rounded(d, (30, 80, 1470, 790), "#fbfdfc", f"#{PALETTE['line']}", 12, 2)
    rounded(d, (52, 108, 88, 140), "#e7f4f2", "#e7f4f2", 16, 1)
    d.text((66, 113), "1", font=small, fill=f"#{PALETTE['accent']}")
    d.text((105, 114), "SOURCE", font=small, fill=f"#{PALETTE['muted']}")
    rounded(d, (185, 104, 1320, 144), "#f4f8f7", "#d9e2df", 7, 1)
    d.text((200, 113), r"Resource\API\global_functions.html", font=mono, fill=f"#{PALETTE['ink']}")
    rounded(d, (1340, 104, 1450, 144), "#edf4ff", "#bfd3ff", 7, 1)
    d.text((1360, 113), "score 0.909", font=mono, fill="#1d4ed8")

    rounded(d, (60, 180, 1440, 735), "#ffffff", "#d9e2df", 10, 1)
    d.text((88, 210), "Section: 全局函数 > tedcore::ac AC仿真指令", font=h2, fill=f"#{PALETTE['ink']}")
    d.text((88, 252), "参数", font=body, fill=f"#{PALETTE['ink']}")

    x0, y0 = 88, 295
    cols = [75, 145, 730, 115, 145]
    headers = ["序号", "参数", "说明", "类型", "默认值"]
    rows = [
        ["1", "start", "扫描起始点。", "float", "100"],
        ["2", "stop", "扫描终点。", "float", "1000000000.0"],
        ["3", "type", "扫描模式。适用于 hspice、hsim、xyce 和 ngspice。", "str", "dec"],
        ["4", "ND", "每十倍频扫描点数。", "int", "10"],
        ["5", "stage", "声明是前仿还是后仿；默认 None 会根据环境变量判断。", "str", "None"],
        ["6", "settings", "添加不常用参数设置，仅限 spectre 仿真器。", "Dict", "{}"],
        ["7", "options", "在网表中添加额外指令，例如 save 指令。", "str", ""],
    ]
    for i, width in enumerate(cols):
        d.rectangle((x0 + sum(cols[:i]), y0, x0 + sum(cols[:i + 1]), y0 + 42), fill="#edf5f3", outline="#b8c7c3")
        d.text((x0 + sum(cols[:i]) + 10, y0 + 10), headers[i], font=small, fill=f"#{PALETTE['ink']}")
    for r, row in enumerate(rows):
        y = y0 + 42 + r * 42
        for i, width in enumerate(cols):
            d.rectangle((x0 + sum(cols[:i]), y, x0 + sum(cols[:i + 1]), y + 42), fill="#ffffff", outline="#cfd8d5")
            d.text((x0 + sum(cols[:i]) + 10, y + 10), row[i], font=small, fill=f"#{PALETTE['ink']}")

    d.text((88, 655), "返回值：Waveform 对象。", font=body, fill=f"#{PALETTE['ink']}")
    img.save(path)


def make_evidence_code(path: Path) -> None:
    img = Image.new("RGB", (1500, 850), "#ffffff")
    d = ImageDraw.Draw(img)
    h1 = _font(31, True)
    h2 = _font(22, True)
    small = _font(16)
    mono = ImageFont.truetype("C:/Windows/Fonts/consola.ttf", 18) if Path("C:/Windows/Fonts/consola.ttf").exists() else small

    d.text((30, 25), "Evidence 证据区示意：命中代码块并展开为完整代码片段", font=h1, fill=f"#{PALETTE['ink']}")
    rounded(d, (30, 80, 1470, 800), "#fbfdfc", f"#{PALETTE['line']}", 12, 2)
    rounded(d, (52, 108, 88, 140), "#e7f4f2", "#e7f4f2", 16, 1)
    d.text((66, 113), "5", font=small, fill=f"#{PALETTE['accent']}")
    d.text((105, 114), "SOURCE", font=small, fill=f"#{PALETTE['muted']}")
    rounded(d, (185, 104, 1320, 144), "#f4f8f7", "#d9e2df", 7, 1)
    d.text((200, 113), r"Resource\optimization\frontend.html", font=mono, fill=f"#{PALETTE['ink']}")
    rounded(d, (60, 180, 1440, 765), "#ffffff", "#d9e2df", 10, 1)
    d.text((88, 210), "Section: 仿真优化底层数据结构 > 样例 2：运行 Xyce AC 仿真并获取结果", font=h2, fill=f"#{PALETTE['ink']}")
    rounded(d, (88, 255, 1415, 735), "#f4f8f7", "#d9e2df", 8, 1)

    code = """from pyted import testbench, module, Vdd, Vss, setenv
from ted_device import Mosfet, VSource, Capacitor, Resistor
from ted_frontend.src.simulator import Simulation

@testbench
def ota_ac_simulation():
    with setenv(simulator="xyce", sensitivity=True):
        (Vdd(), Vss()) >> VSource(dc=1.2).rename("Vdd_src")
        (vp, Vss()) >> VSource(dc=0.6, ac=1).rename("Vp_src")
        (vn, Vss()) >> VSource(dc=0.6).rename("Vn_src")

        ac_raw = Simulation.ac(output="vout", start=1e3, stop=1e10)
        freq = ac_raw.get_sweeps()
        vout_complex = ac_raw.get_signal("vout")

        from ted_frontend.src.torch import T
        gain_db = T.db20(vout_complex)
        phase_deg = T.phase(vout_complex)"""
    y = 280
    for line in code.splitlines():
        d.text((112, y), line, font=mono, fill=f"#{PALETTE['ink']}")
        y += 28

    img.save(path)


def make_assets() -> dict[str, Path]:
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "architecture": IMG_DIR / "architecture.png",
        "ui": IMG_DIR / "ui_overview.png",
        "evidence_table": IMG_DIR / "evidence_table.png",
        "evidence_code": IMG_DIR / "evidence_code.png",
    }
    make_architecture_diagram(assets["architecture"])
    make_ui_overview(assets["ui"])
    make_evidence_table(assets["evidence_table"])
    make_evidence_code(assets["evidence_code"])
    return assets


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.strip("#"))


def set_cell_border(cell, color: str = "d9e2df", size: str = "8") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.strip("#"))


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips: int) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)


def style_paragraph(p, size=10.5, color="17201f", space_after=5, line_spacing=1.15):
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for run in p.runs:
        set_run_font(run, size=size, color=color)


def add_para(doc: Document, text: str = "", size=10.5, color="17201f", bold=False, space_after=5, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(8)
        size = 18
        color = PALETTE["accent"]
    elif level == 2:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(5)
        size = 14
        color = PALETTE["accent2"]
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        size = 12
        color = PALETTE["ink"]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=PALETTE["muted"])
    return p


def add_code_block(doc: Document, code: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 9000)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "f4f8f7")
    set_cell_border(cell, "d9e2df")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.strip("\n").splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=9, color=PALETTE["ink"], name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_note_box(doc: Document, title: str, body: str, fill="eef5f4", accent="0f766e"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, 9000)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "cfd8d5")
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    p = cell.paragraphs[0]
    run = p.add_run(title)
    set_run_font(run, size=10.5, bold=True, color=accent)
    p.add_run().add_break()
    for idx, line in enumerate(body.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=9.5, color=PALETTE["ink"])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc: Document, items: list[str], level: int = 0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        set_run_font(run, size=10, color=PALETTE["ink"])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "edf5f3")
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=9, bold=True, color=PALETTE["ink"])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(text) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            for idx, line in enumerate(str(text).splitlines()):
                if idx:
                    p.add_run().add_break()
                run = p.add_run(line)
                set_run_font(run, size=8.5, color=PALETTE["ink"])
    if widths_cm:
        for row in table.rows:
            for idx, width in enumerate(widths_cm):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def build_doc() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    assets = make_assets()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(f"RagAgent TED 文档问答助手交接文档 | 内网路径：{DEPLOY_PATH}")
    set_run_font(fr, size=8.5, color=PALETTE["muted"])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("RagAgent TED 文档问答助手\n交接文档")
    set_run_font(r, size=28, bold=True, color=PALETTE["ink"])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("范围：仅覆盖 RAGAgent TED 文档问答助手，不包含任务执行、脚本生成、远程运行链路。")
    set_run_font(r, size=12, color=PALETTE["muted"])

    add_table(
        doc,
        ["项目", "说明"],
        [
            ["内网部署路径", DEPLOY_PATH],
            ["前端入口", "http://<内网机IP>:8000/ragagent"],
            ["主要接口", "GET /health；POST /v1/rag/reindex；POST /v1/query/rewrite；POST /v1/rag/ask"],
            ["知识库目录", "Resource/"],
            ["向量索引目录", "workdir/vector_index/"],
            ["问答日志库", "workdir/qa_feedback.db"],
            ["适用人员", "后端维护、部署运维、RAG/检索调优、前端问题排查人员"],
        ],
        [4.0, 12.5],
    )

    add_note_box(
        doc,
        "交接原则",
        "1. 任何 TED 文档更新后都要重新构建向量索引。\n"
        "2. 模型、Embedding、Rerank 的地址和密钥只维护在内网环境变量或 .env 中，不写入交接文档。\n"
        "3. 排查问答质量时先看 Evidence，再判断是文档缺失、切块问题、召回问题还是生成问题。",
    )
    doc.add_page_break()

    add_heading(doc, "1. 文档范围与系统定位", 1)
    add_para(
        doc,
        "RagAgent TED 文档问答助手面向 TED 用户文档、API 文档、仿真/版图/优化示例等资料，提供基于证据的中文问答。"
        "用户可以在前端输入问题，选择是否使用 query rewrite，系统检索相关文档 chunk，经过 rerank 后把证据传给大模型生成答案。",
    )
    add_bullets(
        doc,
        [
            "本文只交接 TED 文档问答助手：Gradio 前端、问答 API、索引构建、检索与证据渲染、历史记录和反馈。",
            "不交接 /v1/tasks/run 的 LangGraph 任务执行、代码生成、preflight、SSH/local runner 等自动执行链路。",
            "若代码中存在共享模块，例如 EmbeddingRetriever、ResourceIndexer、OpenAICompatClient，本文只说明其在问答助手中的使用方式。",
        ],
    )
    doc.add_picture(str(assets["architecture"]), width=Inches(6.55))
    add_caption(doc, "图 1  TED 文档问答助手链路与维护边界")

    add_heading(doc, "2. 部署路径与启动方式", 1)
    add_para(doc, f"内网部署目录固定按当前交接口径记录为：{DEPLOY_PATH}。实际运维时先进入该目录，再执行启动、索引重建和健康检查。")
    add_code_block(
        doc,
        f"""cd {DEPLOY_PATH}

# 首次部署或依赖变更
pip install -r requirements.txt

# 启动服务，默认监听 0.0.0.0:8000
HOST=0.0.0.0 PORT=8000 bash scripts/start_server.sh

# 后台运行示例
nohup bash scripts/start_server.sh > workdir/ragagent.log 2>&1 &""",
    )
    add_table(
        doc,
        ["入口", "用途", "交接说明"],
        [
            ["http://<内网机IP>:8000/ragagent", "前端工作台", "用户使用入口，包含示例提问、可选改写、回答、参考依据、最近对话和评价。"],
            ["http://<内网机IP>:8000/docs", "FastAPI 文档", "用于接口联调；生产内网可视权限策略决定是否开放。"],
            ["http://<内网机IP>:8000/health", "健康检查", "确认模型、Embedding、Rerank 配置是否已注入，以及索引 LATEST 指针。"],
        ],
        [5.0, 3.2, 8.5],
    )

    add_heading(doc, "3. 运行前检查与重建索引", 1)
    add_para(doc, "部署或更新 TED 文档后，先做健康检查，再重建索引。重建成功的基本条件是 doc_count > 0、chunk_count > 0、vector_count == chunk_count。")
    add_code_block(
        doc,
        """curl http://127.0.0.1:8000/health

curl -X POST http://127.0.0.1:8000/v1/rag/reindex

# 或使用脚本
BASE_URL=http://127.0.0.1:8000 bash scripts/reindex.sh""",
    )
    add_bullets(
        doc,
        [
            "重建索引会读取 Resource 下的 .md/.markdown/.html/.htm/.txt 文档。",
            "索引文件保存到 workdir/vector_index，LATEST 文件指向当前生效索引。",
            "reindex 完成后会 invalidate QA cache；如果 RAG_QA_WARMUP_ON_REINDEX=true，还会预热 QA 运行资产。",
            "如果返回 500，优先检查文档目录、Embedding 模型可用性、API base/key、批处理大小和网络连通性。",
        ],
    )

    add_heading(doc, "4. 关键目录和模块", 1)
    add_table(
        doc,
        ["路径", "职责", "维护关注点"],
        [
            ["backend/app.py", "FastAPI 入口；挂载 /ragagent；提供\n/health、/v1/rag/reindex、\n/v1/query/rewrite、/v1/rag/ask。", "新增接口或改路径时同步更新前端和交接文档。"],
            ["backend/ui/workbench.css", "工作台配色、控件、响应式和动效。", "与 UI 模块一起部署；修改后重启服务。"],
            ["backend/ui/gradio_ragagent.py", "Gradio 前端、Evidence HTML 渲染、历史记录、反馈按钮。", "组件交互、证据展示、历史记录异常优先看这里。"],
            ["backend/agents/qa_agent.py", "问答主链路：加载索引、混合召回、重排、拒答、构造提示词、生成答案。", "问答质量、拒答阈值、prompt 策略主要在这里调。"],
            ["backend/agents/\nquery_rewriter.py", "query rewrite：保守型/激进型，保护实体和核心意图。", "检索词扩展过度或不够时调这里。"],
            ["backend/rag/indexer.py", "文档扫描、HTML 正文抽取、表格/代码块处理、切块。", "新增文件类型或切块策略调整在这里。"],
            ["backend/rag/vector_store.py", "Embedding 构建、向量持久化、余弦相似度搜索。", "Embedding 批量失败、索引维度不一致看这里。"],
            ["backend/rag/retriever.py", "BM25-like lexical broad retrieval 和 fallback rerank。", "英文 API 名、参数名精确匹配不足时看这里。"],
            ["backend/rag/evidence.py", "命中代码分片后按 block_id 合并为完整代码块。", "Evidence 中代码不完整时看这里。"],
            ["backend/storage/\nqa_feedback_store.py", "SQLite 历史问答和有用/无用反馈。", "历史记录、反馈写入、数据库迁移看这里。"],
        ],
        [4.0, 8.0, 5.0],
    )

    add_heading(doc, "5. 核心配置项", 1)
    add_para(doc, "配置通过 .env 或环境变量读取。交接时只确认变量名和用途，不在文档中记录真实密钥。")
    add_table(
        doc,
        ["变量", "默认值/示例", "用途", "排查提示"],
        [
            ["RAG_RESOURCE_DIR", "Resource", "TED 原始文档目录。", "无文档或 chunk_count=0 时先检查。"],
            ["RAG_VECTOR_INDEX_DIR", "./workdir/vector_index", "向量索引持久化目录。", "LATEST 缺失会触发重新加载或重建。"],
            ["RAG_QA_FEEDBACK_DB", "./workdir/qa_feedback.db", "历史对话和反馈 SQLite。", "写入失败检查目录权限。"],
            ["OPENAI_API_BASE / OPENAI_API_KEY", "内网模型服务地址/密钥", "回答生成模型。", "answer LLM 报错时检查。"],
            ["MODEL_NAME", "deepseek-v4-pro", "问答和 rewrite 使用的 chat 模型。", "模型下线或改名会导致生成失败。"],
            ["EMBEDDING_API_BASE / EMBEDDING_API_KEY", "内网 embedding 地址/密钥", "构建向量索引和 query 向量。", "reindex 失败优先检查。"],
            ["EMBEDDING_MODEL_TEXT", "text-embedding-v4", "文本 embedding 模型。", "账号不可用时换成实际可用模型。"],
            ["EMBEDDING_BATCH_SIZE", "10", "Embedding 批大小。", "TooLarge 或限流时调小。"],
            ["RERANK_ENABLED", "true", "是否启用二阶段重排。", "不可用会 fallback 到 lexical rerank。"],
            ["RERANK_MODEL_TEXT", "qwen3-rerank", "Rerank 模型。", "证据排序异常或接口失败时检查。"],
            ["RERANK_TOPN_FACTOR", "4", "候选扩展倍数。", "top_n=max(20, top_k*factor)。"],
            ["RAG_QA_TIMING_LOG", "true", "打印问答阶段耗时。", "性能排查建议保留。"],
            ["RAG_QA_WARMUP_ON_REINDEX", "true", "重建索引后预热 QA cache。", "冷启动慢时保留，重建慢时可关闭。"],
        ],
        [3.7, 3.5, 5.0, 4.4],
    )

    add_heading(doc, "6. 文档解析和切块策略", 1)
    add_para(doc, "问答质量高度依赖 ResourceIndexer 的解析和切块。当前实现是自定义结构化切块，不是简单按固定长度直接切全文。")
    add_table(
        doc,
        ["内容类型", "处理策略", "结果"],
        [
            ["HTML / HTM", "使用 HTMLParser；过滤 script/style/nav/aside/header/footer；识别 main/article、main-content、markdown-body 等正文区域；按 h1-h6 标题分段。", "得到带 Section 层级的正文片段，避免导航和模板噪声污染索引。"],
            ["Markdown / TXT", "识别 # 到 ###### 标题；fenced code block 内的 # 不当作标题；无标题文本整体作为 section。", "保留文档层级，使 chunk 带 Section 前缀。"],
            ["表格", "HTML table 收集 tr/td/th 后转 Markdown table；单元格换行转 <br>，竖线转义。", "API 参数、默认值、返回值等结构关系能在 Evidence 中保留。"],
            ["代码块", "HTML pre 与 Markdown fenced code 统一成 fenced block；长代码按行分片，记录 block_id、part_index、part_count。", "向量化时不会因代码过长失真；Evidence 命中后可合并完整代码块。"],
            ["普通长段落", "先段落合并，再按句子或行切；最后才退化为字符滑窗。", "默认 size=1200、overlap=200，短于 50 字符的片段丢弃。"],
        ],
        [3.0, 9.5, 4.5],
    )
    add_note_box(
        doc,
        "切块调优原则",
        "1. 表格和代码是 TED 文档问答的关键证据，尽量保留结构。\n"
        "2. chunk 太小会把 API 名、参数和示例切散；太大会降低检索精度。\n"
        "3. 调整切块策略后必须强制 reindex，并用典型问题检查 Evidence 是否命中正确来源。",
    )

    add_heading(doc, "7. 问答链路", 1)
    add_para(doc, "RagQaAgent.ask(question) 是问答助手的核心入口。链路如下：")
    add_bullets(
        doc,
        [
            "加载运行资产：从 Resource 重新计算 fingerprint；优先读取 workdir/vector_index 中匹配模型和 API base 的缓存索引。",
            "向量召回：EmbeddingRetriever.search 对 query 向量和 chunk 向量做余弦相似度。",
            "关键词召回：HybridRetriever.retrieve_broad 使用 BM25-like 和 token overlap 补充 API 名、参数名、英文术语精确匹配。",
            "候选合并：向量候选和 lexical 候选按 chunk_id 去重合并。",
            "二阶段重排：RERANK_ENABLED=true 时调用 rerank 模型；失败时 warning 记录并降级 lexical rerank。",
            "代码证据展开：expand_code_block_evidence 将命中的代码分片按 block_id 合并为完整代码块。",
            "拒答保护：qa_overlap < 0.15，或 top1_score < 0.70 且 qa_overlap < 0.30 时返回“未检索到相关内容”。",
            "生成答案：模型只能基于 evidence 回答；用法/示例类问题优先给代码块；占位项必须标记“【需自行替换】”。",
        ],
    )

    add_heading(doc, "8. API 交接", 1)
    add_table(
        doc,
        ["接口", "请求", "响应重点", "用途"],
        [
            ["GET /health", "无", "ok、模型配置是否有 key、embedding/rerank 配置、vector_index_latest。", "服务活性和配置检查。"],
            ["POST /v1/rag/reindex", "无 body", "doc_count、fingerprint、chunk_count、vector_count、saved_file、qa_cache_warmed。", "更新文档后重建索引。"],
            ["POST /v1/query/rewrite", '{"query":"...","scene":"qa","mode":"aggressive"}', "original_query、rewritten_query、changed、strategy、warning。", "前端“生成改写建议”按钮对应的代理能力。"],
            ["POST /v1/rag/ask", '{"question":"如何使用 TED 进行 AC 仿真？"}', "status、answer、evidence[]、warning。", "问答主接口。"],
        ],
        [4.1, 4.8, 6.0, 3.0],
    )
    add_code_block(
        doc,
        """curl -X POST http://127.0.0.1:8000/v1/rag/ask \\
  -H "Content-Type: application/json" \\
  -d '{"question":"如何使用 TED 进行 AC 仿真？"}'

# 成功响应核心结构
{
  "status": "answered",
  "answer": "...",
  "evidence": [
    {"source": "Resource/...", "score": 0.909094, "snippet": "Section: ..."}
  ],
  "warning": ""
}""",
    )

    add_heading(doc, "9. 前端工作台", 1)
    add_para(
        doc,
        "前端由 Gradio 挂载在 /ragagent。桌面左侧为“从一个问题开始”，右侧为“回答与发现”；参考依据展示来源、相关度和文本/表格/代码。窄屏采用单栏，提问后自动定位回答。"
        "用户可点击“有帮助 / 需改进”，历史按浏览器标识筛选服务端记录；该标识不提供登录认证。",
    )
    doc.add_picture(str(assets["ui"]), width=Inches(6.55))
    add_caption(doc, "图 2  /ragagent 实际工作台初始页面（2026-09-09）")
    add_table(
        doc,
        ["区域", "功能", "维护点"],
        [
            ["从一个问题开始", "输入或选择示例；点击检索并回答；可选展开改写设置。", "编辑原问题、切换模式或选择新示例会清除旧改写并恢复原始来源。"],
            ["回答与发现", "展示中文状态、提示、Markdown 回答和评价按钮。", "答案为空、not_found 或 error 时看 qa_agent 的 warning 和后端日志。"],
            ["参考依据", "展示来源、相关度和片段；第一条默认展开，支持表格和代码。", "Evidence 不准时先判断是文档、索引、召回、rerank 还是 query rewrite 问题。"],
            ["历史对话", "记录当前浏览器 user_id 下最近问答，支持回填历史问题和证据。", "数据写入 workdir/qa_feedback.db。"],
        ],
        [3.4, 7.2, 6.2],
    )
    doc.add_picture(str(assets["evidence_table"]), width=Inches(6.55))
    add_caption(doc, "图 3  表格类 Evidence 示意")
    doc.add_picture(str(assets["evidence_code"]), width=Inches(6.55))
    add_caption(doc, "图 4  代码类 Evidence 示意")

    add_heading(doc, "10. 数据与持久化", 1)
    add_table(
        doc,
        ["数据", "位置", "内容", "交接建议"],
        [
            ["TED 原始资料", "Resource/", "HTML、Markdown、txt 文档。", "更新前备份；更新后执行 reindex。"],
            ["向量索引", "workdir/vector_index/", "fingerprint.json、LATEST；包含 chunks、vectors、embedding_model、embedding_api_base。", "模型或 embedding base 变化会导致缓存失效，需要重建。"],
            ["问答日志", "workdir/qa_feedback.db", "qa_logs 表：问题、rewrite、答案、证据、反馈、时间。", "定期备份；涉及用户问题，按内网数据规范管理。"],
            ["运行日志", "workdir/*.log 或启动脚本重定向文件", "服务启动、请求错误、RAG-QA-Timing。", "排查性能和模型异常时保留。"],
        ],
        [3.0, 5.0, 5.2, 4.0],
    )

    add_heading(doc, "11. 日常维护 SOP", 1)
    add_heading(doc, "11.1 更新 TED 文档", 2)
    add_bullets(
        doc,
        [
            "把新文档放入 Resource，保持目录结构稳定，避免混入无关 HTML 模板或临时文件。",
            "执行 /v1/rag/reindex 或 scripts/reindex.sh。",
            "检查返回 doc_count、chunk_count、vector_count，并确认 vector_count == chunk_count。",
            "用 3-5 个典型问题在 /ragagent 检查 Evidence 来源是否正确。",
        ],
    )
    add_heading(doc, "11.2 修改模型或 Embedding", 2)
    add_bullets(
        doc,
        [
            "修改 .env 中对应 API base、key、model 名称。",
            "重启服务，使配置重新加载。",
            "执行 /health 确认 key_set 和 model 字段。",
            "若修改的是 embedding 模型或 base，必须重建索引。",
        ],
    )
    add_heading(doc, "11.3 例行验收问题", 2)
    add_table(
        doc,
        ["问题", "预期检查点"],
        [
            ["如何使用 TED 进行 AC 仿真？", "应命中 ac 函数、参数表、raw.get_signal/get_sweeps 等证据。"],
            ["AC 仿真结果怎么获取幅度和相位？", "应出现 signal_type 或 T.db20/T.phase 等相关证据。"],
            ["如何在 TED 中做 Monte Carlo 仿真？", "rewrite 应补充 Monte Carlo 术语，Evidence 应来自仿真相关文档。"],
            ["某个 API 参数默认值是什么？", "应优先命中 API 表格证据，而不是泛泛解释。"],
        ],
        [6.0, 10.5],
    )

    add_heading(doc, "12. 常见问题排查", 1)
    add_table(
        doc,
        ["现象", "优先排查", "处理建议"],
        [
            ["/ragagent 无法访问", "gradio 是否安装；backend.app 是否成功 import mount_ragagent_ui；端口是否监听。", "执行 pip install -r requirements.txt；重启服务；看启动日志。"],
            ["/health ok 但问答失败", "chat model、embedding、rerank 是否可用；API key 是否注入。", "看 /health 的 key_set；用最小 curl 请求验证模型服务。"],
            ["reindex 返回 vector_count 不等于 chunk_count", "Embedding 返回数量异常或中途中断。", "降低 EMBEDDING_BATCH_SIZE；检查 embedding 服务日志；清理失败索引后重建。"],
            ["Evidence 来源不相关", "Resource 是否已更新但未 reindex；query rewrite 是否漂移；rerank 是否失败。", "先关闭 rewrite 试问；查看 warning；必要时调 query_rewriter 或 retriever。"],
            ["答案编造 API 或参数", "证据不足但模型未拒答；prompt 或拒答阈值偏松。", "检查 qa_agent 的 _should_abstain、prompt 规则和 Evidence 内容。"],
            ["代码 Evidence 不完整", "代码块分片合并逻辑或 block_id 元数据。", "检查 backend/rag/evidence.py 和 indexer 的 code block 分片。"],
            ["历史记录不显示", "qa_feedback.db 权限；user_id 生成；SQLite WAL 文件。", "检查 workdir 目录可写；查看 qa_feedback_store 写入异常。"],
        ],
        [4.3, 6.2, 6.2],
    )

    add_heading(doc, "13. 交接清单", 1)
    add_table(
        doc,
        ["检查项", "完成标准"],
        [
            ["代码目录确认", "能进入部署目录，且 requirements.txt、backend、Resource、scripts、workdir 存在。"],
            ["服务启动", "scripts/start_server.sh 能正常启动，/health 返回 ok=true。"],
            ["前端访问", "/ragagent 能打开；初始显示等待提问；示例填入和可选改写展开正常。"],
            ["索引重建", "/v1/rag/reindex 返回 chunk_count > 0 且 vector_count == chunk_count。"],
            ["问答验收", "典型 TED 问题能返回 answered，Evidence 来源正确。"],
            ["反馈日志", "点击有用/无用后 qa_feedback.db 中记录可查询。"],
            ["配置交接", ".env 中模型、Embedding、Rerank 配置由内网运维安全交接，不写入文档。"],
            ["回滚方案", "保留上一版 Resource、vector_index 和 qa_feedback.db 备份。"],
        ],
        [5.0, 11.5],
    )

    add_note_box(
        doc,
        "最后说明",
        "后续如果要继续提升问答质量，优先基于 qa_feedback.db 中的 bad case 观察 Evidence："
        "如果 Evidence 不对，调文档解析、切块、召回和 rerank；如果 Evidence 对但答案不稳，调 qa_agent 的 prompt 和拒答策略。",
        fill="f8fafc",
        accent=PALETTE["accent2"],
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_doc()
    print(path)
